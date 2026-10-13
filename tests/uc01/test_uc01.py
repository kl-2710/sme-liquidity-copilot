"""Kiểm thử UC-01 theo các giá trị tham chiếu trong workbook TV1."""
from datetime import date
import os
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from src.core import data_loader
from src.core.data_loader import DataValidationError, get_snapshot
from src.uc01_working_capital.working_capital import _buyers, _suppliers, calculate_working_capital
from src.agent.tools import get_working_capital, simulate_scenario


class UC01Tests(unittest.TestCase):
    def test_f01_reconciles_to_workbook(self):
        result = calculate_working_capital(get_snapshot("F01"))
        metrics = {m.id: m.raw_value for m in result.result_metrics}
        self.assertEqual(metrics["ar_vnd"], 4_912_255_000)
        self.assertEqual(metrics["inv_vnd"], 3_450_000_000)
        self.assertEqual(metrics["ap_vnd"], 1_951_577_000)
        self.assertAlmostEqual(metrics["dso_days"], 49.98858558235239, places=8)
        self.assertAlmostEqual(metrics["dio_days"], 43.886925795053, places=8)
        self.assertAlmostEqual(metrics["dpo_days"], 24.825714487632506, places=8)
        self.assertAlmostEqual(metrics["ccc_days"], 69.04979688977288, places=8)
        self.assertTrue(result.reconciliation.matches_ar)
        self.assertTrue(result.reconciliation.matches_inv)
        self.assertTrue(result.reconciliation.matches_ap)
        self.assertTrue(result.reconciliation.matches_ccc)

    def test_extended_metrics_match_reference(self):
        result = calculate_working_capital(get_snapshot("F01"))
        metrics = {m.id: m.raw_value for m in result.result_metrics}
        self.assertEqual(metrics["ar_overdue_vnd"], 871_565_000)
        self.assertAlmostEqual(metrics["ar_overdue_share"], 0.17742666046449135)
        self.assertEqual(metrics["slow_inventory_vnd"], 430_000_000)
        self.assertEqual(metrics["ap_overdue_vnd"], 341_170_000)
        self.assertAlmostEqual(metrics["gross_margin"], 0.20002966957098298)
        self.assertAlmostEqual(metrics["cash_per_1_day_dpo_pur_vnd"], 47_093_711.11111111)
        self.assertEqual(sum(s["saving_vnd"] for s in result.model_dump(mode="json")["drivers"]["early_payment_discounts"]), 4_501_240)

    def test_all_sample_firms_reconcile(self):
        expected = {
            "F01": (4_912_255_000, 3_450_000_000, 1_951_577_000, 69.04979688977288),
            "F02": (4_807_643_000, 8_050_000_000, 2_317_750_000, 85.99180048178746),
            "F03": (3_881_055_000, 1_860_000_000, 888_463_000, 75.3305891013971),
        }
        for firm_id, values in expected.items():
            with self.subTest(firm_id=firm_id):
                result = calculate_working_capital(get_snapshot(firm_id))
                metrics = {m.id: m.raw_value for m in result.result_metrics}
                self.assertEqual(metrics["ar_vnd"], values[0])
                self.assertEqual(metrics["inv_vnd"], values[1])
                self.assertEqual(metrics["ap_vnd"], values[2])
                self.assertAlmostEqual(metrics["ccc_days"], values[3], places=8)
                self.assertTrue(all((result.reconciliation.matches_ar, result.reconciliation.matches_inv,
                                     result.reconciliation.matches_ap, result.reconciliation.matches_ccc)))

    def test_hgc_is_excluded_and_reported(self):
        result = calculate_working_capital(get_snapshot("F01"))
        self.assertEqual(result.inputs.inv_vnd, 3_450_000_000)
        self.assertEqual(result.inputs.excluded_inventory_vnd, 2_600_000_000)
        self.assertEqual(len(result.drivers.inventory_exclusions), 1)

    def test_output_contract_and_synthetic_warning(self):
        result = get_working_capital("F01")
        payload = result["result"]
        self.assertTrue(payload["meta"]["data_is_synthetic"])
        self.assertEqual(payload["formula_version"], "tv1.formulas.v4")
        self.assertIn("metrics", payload["result"])
        self.assertGreaterEqual(len(payload["result"]["metrics"]), 20)
        self.assertTrue(payload["unsupported_metrics"])
        self.assertTrue(payload["data_quality"])
        # Cửa sổ đến hạn cộng dồn (net-of-VAT): today <= 7d <= 30d <= AP của NCC
        for s in payload["drivers"]["suppliers"]:
            self.assertEqual(s["due_basis"], "net_of_VAT")
            self.assertLessEqual(s["due_today_vnd"], s["due_7d_vnd"] + 1e-6)
            self.assertLessEqual(s["due_7d_vnd"], s["due_30d_vnd"] + 1e-6)
            self.assertLessEqual(s["due_30d_vnd"], s["ap_vnd"] + 1e-6)
        self.assertEqual(payload["reconciliation"]["details"]["aging_vnd"], 0)

    def test_scenario_uses_raw_uc01_values_and_marks_demo_forecast(self):
        result = simulate_scenario("F01", {"dso_change_days": -1, "dio_change_days": 0, "dpo_change_days": 0})
        self.assertEqual(result["result"]["status"], "provisional_demo")
        self.assertAlmostEqual(result["result"]["cash_released_vnd"], 98_267_533, delta=1)
        self.assertEqual(result["result"]["before"]["ccc_days"], 69.0)
        self.assertEqual(result["result"]["after"]["ccc_days"], 68.0)
        self.assertIn("demo", result["result"]["note"].lower())

    def test_partial_payment_residual_is_not_face_value(self):
        as_of = date(2026, 9, 30)
        sales = [{"firm_id": "F01", "invoice_id": "S-1", "buyer_id": "B1", "buyer_name": "Buyer",
            "issue_date": date(2026, 9, 1), "due_date": as_of, "invoice_amount": 100.0,
            "paid_amount": 40.0, "credit_amount": 0.0, "paid_date": date(2026, 9, 20), "source_row_id": "S-1"}]
        events = [{"invoice_id": "S-1", "amount_vnd": 40.0, "paid_date": date(2026, 9, 20)}]
        buyers, ar = _buyers(sales, as_of, 1000.0, 90, events)
        self.assertEqual(ar, 60.0)
        self.assertEqual(buyers[0].ar_vnd, 60.0)
        self.assertEqual(buyers[0].overdue_ar_vnd, 0.0)

    def test_due_date_boundary_and_aging(self):
        as_of = date(2026, 9, 30)
        sales = [
            {"firm_id": "F01", "invoice_id": "S-DUE", "buyer_id": "B1", "buyer_name": "Buyer 1",
             "issue_date": date(2026, 9, 1), "due_date": as_of, "invoice_amount": 100.0,
             "paid_amount": 0.0, "credit_amount": 0.0, "paid_date": None, "source_row_id": "S-DUE"},
            {"firm_id": "F01", "invoice_id": "S-LATE", "buyer_id": "B2", "buyer_name": "Buyer 2",
             "issue_date": date(2026, 8, 1), "due_date": date(2026, 9, 29), "invoice_amount": 200.0,
             "paid_amount": 0.0, "credit_amount": 0.0, "paid_date": None, "source_row_id": "S-LATE"},
        ]
        buyers, ar = _buyers(sales, as_of, 1000.0, 90, [])
        by_id = {b.buyer_id: b for b in buyers}
        self.assertEqual(ar, 300.0)
        self.assertEqual(by_id["B1"].overdue_ar_vnd, 0.0)
        self.assertEqual(by_id["B1"].aging_vnd["not_due"], 100.0)
        self.assertEqual(by_id["B2"].overdue_ar_vnd, 200.0)
        self.assertEqual(by_id["B2"].aging_vnd["1_30"], 200.0)

    def test_payment_after_as_of_does_not_reduce_ar(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = data_loader.WORKBOOK_PATH
            target = os.path.join(tmp, "test-workbook.xlsx")
            shutil.copy2(source, target)
            workbook = openpyxl.load_workbook(target)
            sheet = workbook["D_sales"]
            header = [cell.value for cell in sheet[2]]
            paid_col = header.index("paid_date") + 1
            invoice_col = header.index("invoice_id") + 1
            for row in range(3, sheet.max_row + 1):
                if sheet.cell(row, invoice_col).value == "S-F01-0001":
                    sheet.cell(row, paid_col).value = date(2026, 10, 5)
                    break
            workbook.save(target)
            workbook.close()
            with patch.object(data_loader, "WORKBOOK_PATH", type(source)(target)):
                snapshot = get_snapshot("F01")
            invoice = next(s for s in snapshot["sales"] if s["invoice_id"] == "S-F01-0001")
            self.assertEqual(invoice["paid_amount"], 0.0)
            self.assertTrue(invoice["is_open"])

    def test_duplicate_invoice_id_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = data_loader.WORKBOOK_PATH
            target = os.path.join(tmp, "duplicate-workbook.xlsx")
            shutil.copy2(source, target)
            workbook = openpyxl.load_workbook(target)
            sheet = workbook["D_sales"]
            sheet.append([cell.value for cell in sheet[3]])
            workbook.save(target)
            workbook.close()
            with patch.object(data_loader, "WORKBOOK_PATH", type(source)(target)):
                with self.assertRaisesRegex(DataValidationError, "invoice_id trùng lặp"):
                    get_snapshot("F01")

    def test_early_payment_discount_rate_2_10_net_30(self):
        as_of = date(2026, 10, 1)
        purchases = [{"firm_id": "F01", "purchase_id": "P-1", "supplier_id": "SUP1",
            "supplier_name": "Supplier", "purchase_date": date(2026, 9, 25),
            "due_date": date(2026, 10, 25), "amount_vnd": 1000.0,
            "paid_amount_vnd": 0.0, "credit_amount_vnd": 0.0, "paid_date": None,
            "early_discount_rate": 0.02, "days_early": 20,
            "discount_deadline": date(2026, 10, 5), "source_row_id": "P-1"}]
        suppliers, ap = _suppliers(purchases, as_of, 1000.0, 90, [])
        opportunity = suppliers[0].discount_opportunities[0]
        self.assertEqual(ap, 1000.0)
        self.assertEqual(opportunity.days_early, 20)
        self.assertAlmostEqual(opportunity.simple_annual_rate_pct, 37.244897959183675, places=8)

    def test_unknown_firm_is_not_silently_zero(self):
        with self.assertRaises(DataValidationError):
            get_snapshot("DOES_NOT_EXIST")


if __name__ == "__main__":
    unittest.main(verbosity=2)